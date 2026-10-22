from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADERS = ["Code", "Name", "Type", "Born year", "Color", "Health Status", "Breed", "Pedigree"]


def _workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pet Manager"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6FED")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=10)
        sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 36)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _rows(pets, animal_type):
    return [
        [pet.code, pet.name, animal_type, pet.born_year, pet.color, pet.health_status,
         pet.breed, "Yes" if animal_type == "Dog" and pet.pedigree else "No" if animal_type == "Dog" else ""]
        for pet in pets
    ]


def export_pet_list(dogs, cats):
    return _workbook(_rows(dogs, "Dog") + _rows(cats, "Cat"))


def export_dog_list(dogs):
    return _workbook(_rows(dogs, "Dog"))


def export_cat_list(cats):
    return _workbook(_rows(cats, "Cat"))
